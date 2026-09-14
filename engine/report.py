"""
Excel export of a lifetime run. Five tabs:

    Summary       inputs, headline numbers, weather
    Years         one row per simulated year
    Daily         loading / rh_eq / dew point / pane min / film, per day
    Year 1 hourly the animation trace, 8,760 rows
    Assumptions   every estimate the run rests on

US customary at the boundary, matching the UI.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from engine import psychro

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
KEY = Font(bold=True)


def _header(ws, row, cols):
    for c, name in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=c, value=name)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(horizontal="center")


def _autowidth(ws, minimum=10, maximum=48):
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minimum, min(maximum, width + 2))


def build_workbook(result, echo: dict, headline: dict, weather_desc: dict, assumptions: list[str]) -> Workbook:
    wb = Workbook()

    # Summary ---------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "INOVUES Desiccant Lifetime Simulator"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Doc ID ANLY-003. Every number below is a model output on estimated inputs; see Assumptions."
    r = 4
    _header(ws, r, ["Headline", "Value"]); r += 1
    rows = [
        ("Exhausted after (hours)", headline["exhausted_hour"]),
        ("Exhausted after (days)", headline["exhausted_days"]),
        ("Exhausted after (years)", headline["exhausted_years"]),
        ("First condensation (hours)", headline["first_condensation_hour"]),
        ("First condensation (days)", headline["first_condensation_days"]),
        ("Hours per gram", headline["hours_per_gram"]),
        ("Years simulated", headline["years_run"]),
        ("Run capped at max years", headline["capped"]),
        ("Final loading (kg water / kg desiccant)", headline["final_loading"]),
        ("Final loading (% of 25 degC capacity)", headline["final_loading_pct_of_max"]),
        ("Final equilibrium RH of desiccant (%)", headline["final_rh_eq_pct"]),
        ("Total water into desiccant (g)", headline["total_water_into_desiccant_g"]),
    ]
    c = result.contributions()
    pct = lambda v: "n/a" if v is None else round(v, 1)
    rows += [
        ("Net water over desiccant life: outdoor path (g)", round(c["outdoor_g"], 2)),
        ("Net water over desiccant life: room path (g)", round(c["room_g"], 2)),
        ("Net water over desiccant life: sealant diffusion (g)", round(c["diffusion_g"], 2)),
        ("Share of net total: outdoor (%)", pct(c["outdoor_pct"])),
        ("Share of net total: room (%)", pct(c["room_pct"])),
        ("Share of net total: diffusion (%)", pct(c["diffusion_pct"])),
        ("Note", "Negative = that path removed water (its air was drier than the cavity)."),
    ]
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = KEY
        ws.cell(row=r, column=2, value="n/a" if v is None else v)
        r += 1
    r += 1
    _header(ws, r, ["Input", "Value"]); r += 1
    for k, v in echo.items():
        ws.cell(row=r, column=1, value=k).font = KEY
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    _header(ws, r, ["Weather", "Value"]); r += 1
    for k, v in weather_desc.items():
        ws.cell(row=r, column=1, value=k).font = KEY
        ws.cell(row=r, column=2, value=v)
        r += 1
    _autowidth(ws)

    # Years -----------------------------------------------------------
    ws = wb.create_sheet("Years")
    _header(ws, 1, ["Year", "Condensed g/m2", "Hours condensing", "Hours visible",
                    "Water into desiccant g", "Loading end kg/kg", "Desiccant RH_eq end %",
                    "Mean cavity dew point degF",
                    "Net outdoor g", "Net room g", "Net diffusion g",
                    "Heating season net outdoor g", "Heating season net room g"])
    for i, y in enumerate(result.years, start=2):
        ws.append([y.year, round(y.condensed_kg_per_m2 * 1000, 2), y.hours_condensing, y.hours_visible,
                   round(y.water_into_desiccant_g, 2), round(y.loading_end, 4),
                   round(100 * y.rh_eq_end, 2), round(psychro.c_to_f(y.mean_cavity_dew_point_c), 1),
                   round(y.net_outdoor_g, 2), round(y.net_room_g, 2), round(y.net_diffusion_g, 2),
                   round(y.heating_outdoor_g, 2), round(y.heating_room_g, 2)])
    _autowidth(ws)

    # Daily -----------------------------------------------------------
    ws = wb.create_sheet("Daily")
    _header(ws, 1, ["Day", "Loading kg/kg", "Desiccant RH_eq %", "Cavity dew point degF",
                    "Pane min degF", "Film max um"])
    for i in range(len(result.daily_loading)):
        ws.append([i + 1, round(result.daily_loading[i], 5), round(100 * result.daily_rh_eq[i], 3),
                   round(psychro.c_to_f(result.daily_cav_dew_c[i]), 2),
                   round(psychro.c_to_f(result.daily_pane_min_c[i]), 2),
                   round(1000 * result.daily_film_max_kg[i], 3)])
    _autowidth(ws)

    # Year 1 hourly ---------------------------------------------------
    ws = wb.create_sheet("Year 1 hourly")
    y1 = result.year1
    if y1:
        _header(ws, 1, ["Hour", "Outdoor degF", "Cold pane degF", "Cavity air degF", "Cavity RH %",
                        "Cavity dew point degF", "Film um", "Loading kg/kg", "Desiccant RH_eq %",
                        "Uptake g", "Condensed g", "Outdoor ACH"])
        n = len(y1["loading"])
        for i in range(n):
            ws.append([i, round(psychro.c_to_f(y1["t_out_c"][i]), 1), round(psychro.c_to_f(y1["t_cold_c"][i]), 1),
                       round(psychro.c_to_f(y1["t_air_c"][i]), 1), round(100 * y1["rh_cav"][i], 2),
                       round(psychro.c_to_f(y1["dew_cav_c"][i]), 1), round(1000 * y1["film_kg"][i], 2),
                       round(y1["loading"][i], 5), round(100 * y1["rh_eq"][i], 2),
                       round(y1["uptake_g"][i], 4), round(y1["condensed_g"][i], 4), round(y1["ach_out"][i], 3)])
        ws.freeze_panes = "A2"
    _autowidth(ws)

    # Assumptions -----------------------------------------------------
    ws = wb.create_sheet("Assumptions")
    _header(ws, 1, ["#", "Assumption"])
    for i, a in enumerate(assumptions, start=1):
        ws.append([i, a])
    _autowidth(ws, maximum=110)
    return wb


def workbook_bytes(result, echo, headline, weather_desc, assumptions) -> bytes:
    buf = BytesIO()
    build_workbook(result, echo, headline, weather_desc, assumptions).save(buf)
    return buf.getvalue()
