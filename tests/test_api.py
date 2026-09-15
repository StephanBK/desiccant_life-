"""API wiring: presets, lifetime, sweeps, export. Weather is the 277 Park
fixture served in place of the live NSRDB fetch."""

from __future__ import annotations

import gzip
import io
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from engine.weather import Location, parse_nsrdb_csv

FIXTURE = Path(__file__).parent / "fixtures" / "nsrdb_277park_tmy_v4.csv.gz"


@pytest.fixture(scope="module")
def year():
    return parse_nsrdb_csv(gzip.open(FIXTURE, "rt").read())


@pytest.fixture
def client(monkeypatch, year):
    os.environ.setdefault("MAPBOX_TOKEN", "pk.test")
    os.environ.setdefault("NREL_API_KEY", "test")
    import app as flask_app
    loc = Location("x", "277 PARK AVE, NEW YORK, NY 10172", 40.7568, -73.9742)
    monkeypatch.setattr(flask_app, "get_weather_for_address", lambda *a, **k: (loc, year, False))
    return flask_app.app.test_client()


def test_presets_shape(client):
    d = client.get("/api/presets").get_json()
    assert d["defaults"]["al_out"] == "wet_sealed" and d["defaults"]["al_in"] == "wet_sealed"
    assert [p["key"] for p in d["al_out"]][:2] == ["hermetic", "wet_sealed"] and [p["key"] for p in d["al_in"]][:2] == ["hermetic", "wet_sealed"]
    assert d["leakage_reference"]["cp_table"][0] == {"angle_deg": 0.0, "cp": 0.6}
    assert d["defaults"]["p_occ_pa"] == 5.0 and d["defaults"]["p_unocc_pa"] == 0.0
    assert {v["key"] for v in d["sealants"]} == {"silicone", "pib", "none"}
    assert d["leakage_reference"]["aerc_url"].startswith("https://aercenergyrating.org")
    assert d["desiccants"][0]["key"] == "ms3a"
    assert "desiccant_grams" in d["sweepable"]


def test_lifetime_defaults(client):
    d = client.get("/api/lifetime").get_json()
    h = d["headline"]
    assert h["exhausted_hour"] is not None and h["hours_per_gram"] > 0
    assert d["inputs"]["al_out"] == 0.005 and d["inputs"]["al_in"] == 0.005 and d["inputs"]["dp_pa"] == 3.0
    assert d["inputs"]["ach_out"] == pytest.approx(0.005 * 18.29 * (3 / 75) ** 0.65 / (0.6 * 0.0254), rel=1e-3)
    assert d["inputs"]["sealant_in"] == "silicone" and 0.02 < d["inputs"]["diffusion_g_per_day"] < 0.05
    assert d["inputs"]["capacity_g"] == pytest.approx(210.0)
    e = client.get("/api/lifetime?al_out=0&al_in=0&wind_scaling=0&trace=0").get_json()["headline"]
    f = client.get("/api/lifetime?al_out=0&al_in=0&wind_scaling=0&sealant_out=pib&sealant_in=pib&trace=0").get_json()["headline"]
    assert e["exhausted_hour"] is not None and 5 < e["exhausted_years"] < 9        # silicone diffusion floor
    assert f["exhausted_hour"] is None                                            # PIB: beyond 20 years
    assert len(d["year1"]["loading_pct"]) == 8760
    assert len(d["daily"]["loading"]) == 365
    assert d["year1"]["loading_pct"][-1] <= 100.0


def test_lifetime_units_convert_at_boundary(client):
    d = client.get("/api/lifetime?t_in=68&rh_in=40&width_in=48&height_in=72&offset_in=1.0").get_json()
    assert d["inputs"]["t_in_f"] == 68.0 and d["inputs"]["rh_in_pct"] == 40.0
    assert d["inputs"]["offset_in"] == pytest.approx(1.0)


def test_lifetime_more_grams_lasts_longer(client):
    a = client.get("/api/lifetime?grams=20&trace=0").get_json()["headline"]
    b = client.get("/api/lifetime?grams=200&trace=0").get_json()["headline"]
    assert b["exhausted_hour"] > a["exhausted_hour"]
    assert "year1" not in client.get("/api/lifetime?trace=0").get_json()


def test_lifetime_preset_names_and_numbers(client):
    a = client.get("/api/lifetime?al_out=baseline&trace=0").get_json()["inputs"]
    b = client.get("/api/lifetime?al_out=2.0&trace=0").get_json()["inputs"]
    assert a["al_out"] == b["al_out"] == 2.0 and a["ach_out"] == b["ach_out"]


def test_raw_ach_override_and_offset_dilution(client):
    raw = client.get("/api/lifetime?ach_out=1.0&ach_in=0.5&trace=0").get_json()["inputs"]
    assert raw["al_out"] is None and raw["ach_out"] == 1.0 and raw["ach_in"] == 0.5
    thin = client.get("/api/lifetime?offset_in=0.6&trace=0").get_json()["inputs"]["ach_in"]
    deep = client.get("/api/lifetime?offset_in=2.4&trace=0").get_json()["inputs"]["ach_in"]
    assert thin == pytest.approx(4 * deep, rel=1e-3)          # same crack flow, 4x the air


@pytest.mark.parametrize("q", ["f_cold=1.5", "grams=-1", "al_in=lots", "ach_in=-1", "dp_pa=-1", "desiccant=silica", "sealant_in=tar", "bead_depth_in=0",
                               "orientation=up", "offset_in=0"])
def test_lifetime_rejects_bad_input(client, q):
    assert client.get("/api/lifetime?" + q).status_code == 400


def test_sweep_1d(client):
    d = client.get("/api/sweep?x_key=desiccant_grams&x_values=10,50,200&trace=0").get_json()
    assert d["mode"] == "1d" and [p["x"] for p in d["points"]] == [10, 50, 200]
    hrs = [p["exhausted_hour"] for p in d["points"]]
    assert hrs == sorted(hrs)


def test_sweep_2d_user_units(client):
    d = client.get("/api/sweep?x_key=desiccant_grams&x_values=20,100&y_key=t_room_c&y_values=65,75").get_json()
    e = client.get("/api/sweep?x_key=al_in&x_values=0.01,0.06,0.3&trace=0").get_json()
    hrs = [p["exhausted_hour"] for p in e["points"]]
    assert hrs == sorted(hrs, reverse=True)                     # tighter retrofit, longer life
    assert d["mode"] == "2d" and d["y_values"] == [65, 75]
    assert d["grid"][1][1]["y"] == 75 and d["grid"][0][1]["x"] == 100


def test_sweep_rejects(client):
    assert client.get("/api/sweep?x_key=moon&x_values=1").status_code == 400
    assert client.get("/api/sweep?x_key=ach_in&x_values=1&y_key=ach_in&y_values=2").status_code == 400
    big = ",".join(str(i) for i in range(41))
    assert client.get(f"/api/sweep?x_key=desiccant_grams&x_values={big}").status_code == 400


def test_export_workbook(client):
    r = client.get("/api/export.xlsx?grams=50")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Summary", "Years", "Daily", "Year 1 hourly", "Assumptions"]
    assert wb["Year 1 hourly"].max_row == 8761
    assert wb["Summary"]["A5"].value == "Exhausted after (hours)"


def test_health(client):
    assert client.get("/api/health").get_json()["ok"] is True
